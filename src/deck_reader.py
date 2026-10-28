"""讀一份 pitch deck，草擬 16 個子分數、每個說法的證據和要查核的地方。

用法
    python deck_reader.py deck.pdf            # 需要 ANTHROPIC_API_KEY
    python deck_reader.py deck.pdf --dry-run  # 只抽文字、存 prompt，不呼叫 API

輸出  <deck 名>_draft_scores.json，可直接餵給 deck_scorer.py
規則  所有分數一律標 claimed，deck 是創辦人自己說的，驗證是人的工作
"""
import argparse
import json
import os
from pathlib import Path

RUBRIC_PAGE = Path(__file__).parent.parent / 'RUBRIC.md'
RUBRIC = json.loads((Path(__file__).parent / 'rubric.json').read_text())
MODEL = 'claude-opus-5-5'


def extract_deck_text(deck_path):
    """把 PDF 或 PPTX 轉成逐頁文字。

    :param deck_path: deck 檔案路徑
    :type deck_path: Path
    :return: 每頁前面標 [page N] 的全文
    :rtype: str
    """
    pages = []
    if deck_path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        for page in PdfReader(deck_path).pages:
            pages.append(page.extract_text() or '')
    elif deck_path.suffix.lower() == '.pptx':
        from pptx import Presentation
        for slide in Presentation(deck_path).slides:
            slide_text = [shape.text_frame.text for shape in slide.shapes
                          if shape.has_text_frame]
            pages.append('\n'.join(slide_text))
    else:
        pages.append(deck_path.read_text(encoding='utf-8'))  # 純文字測試用
    return '\n\n'.join(f'[page {number}]\n{text}'
                       for number, text in enumerate(pages, start=1))


def build_prompt(deck_text):
    """組合給模型的指示，rubric 全文一起送，分數定義才一致。"""
    criteria = RUBRIC['team'] + RUBRIC['upside'] + RUBRIC['thesis']
    return f"""You are helping a venture scout pre-read a pitch deck.
Score each sub-criterion 1-5 using ONLY the rubric below and ONLY what the
deck says. Everything in a deck is a claim, so never mark anything verified.

RUBRIC
{RUBRIC_PAGE.read_text(encoding='utf-8')}

Return JSON only, in this shape:
{{
  "company": "...",
  "core_risk_axis": "scientific | distribution | infrastructure",
  "red_flags": ["..."],
  "scores": {{
    "<criterion>": {{
      "score": 1-5,
      "status": "claimed",
      "deck_evidence": "short quote from the deck",
      "deck_page": N,
      "how_to_verify": "the exact record to check, e.g. Scholar profile",
      "missing": true or false
    }}
  }},
  "discovery_questions": ["questions for the call about what the deck
                           does not say"]
}}
Criteria keys, use exactly these: {criteria}
If the deck says nothing on a criterion, score 1 and set missing true.

DECK
{deck_text}"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('deck', type=Path)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()

    prompt = build_prompt(extract_deck_text(args.deck))
    out_path = args.deck.with_name(args.deck.stem + '_draft_scores.json')

    if args.dry_run or not os.environ.get('ANTHROPIC_API_KEY'):
        prompt_path = args.deck.with_name(args.deck.stem + '_prompt.txt')
        prompt_path.write_text(prompt, encoding='utf-8')
        print('no API call, prompt saved to', prompt_path)
        return

    import anthropic
    client = anthropic.Anthropic()
    response = client.messages.create(
        model=MODEL,
        max_tokens=8000,
        messages=[{'role': 'user', 'content': prompt}],
    )
    reply = response.content[0].text
    json_text = reply[reply.find('{'):reply.rfind('}') + 1]  # 去掉多餘文字
    draft = json.loads(json_text)
    out_path.write_text(json.dumps(draft, indent=2, ensure_ascii=False),
                        encoding='utf-8')
    print('draft scores saved to', out_path)


if __name__ == '__main__':
    main()
